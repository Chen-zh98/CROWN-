import io
import re
import datetime

import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

st.set_page_config(
    page_title="万商月度下单金额对比",
    page_icon="📊",
    layout="wide",
)


def parse_year_month(filename: str, default_year: int):
    """
    从文件名识别年月。
    支持：
      6月万商.xls
      7月万商.xls
      2025年6月万商.xls
      2025-06万商.xls
      202506万商.xls
    """
    name = filename

    # 2025年6月 / 2025-06 / 2025_06 / 202506
    m = re.search(r"(20\d{2})\s*[年\-_/]?\s*(\d{1,2})\s*月?", name)
    if m:
        y, mo = int(m.group(1)), int(m.group(2))
        if 1 <= mo <= 12:
            return y, mo

    # 6月 / 7月
    m = re.search(r"(\d{1,2})\s*月", name)
    if m:
        mo = int(m.group(1))
        if 1 <= mo <= 12:
            return default_year, mo

    raise ValueError(
        f"无法从文件名识别月份：{filename}。"
        "请把文件名改成类似“6月万商.xls”或“2025年6月万商.xls”。"
    )


def read_one_file(uploaded_file, default_year: int):
    """
    读取一个上传文件，返回明细 DataFrame 和月份 Period。
    """
    suffix = uploaded_file.name.lower().split(".")[-1]
    engine = "xlrd" if suffix == "xls" else "openpyxl"

    year, month = parse_year_month(uploaded_file.name, default_year)
    period = pd.Period(year=year, month=month, freq="M")

    sheets = pd.read_excel(uploaded_file, sheet_name=None, engine=engine)
    frames = []

    for sheet_name, df in sheets.items():
        if df is None or df.empty:
            continue

        df = df.copy()
        df.columns = [str(c).strip().replace("\ufeff", "") for c in df.columns]

        # 兼容“下单帐号 / 下单账号”
        account_col = None
        for c in ["下单帐号", "下单账号"]:
            if c in df.columns:
                account_col = c
                break

        if account_col is None or "订单金额" not in df.columns:
            continue

        out = pd.DataFrame(
            {
                "下单帐号": df[account_col].astype(str).str.strip(),
                "订单金额": pd.to_numeric(
                    df["订单金额"]
                    .astype(str)
                    .str.replace(",", "", regex=False)
                    .str.replace("，", "", regex=False)
                    .str.strip(),
                    errors="coerce",
                ),
            }
        )

        out = out[
            out["下单帐号"].notna()
            & (out["下单帐号"] != "")
            & (out["下单帐号"].str.lower() != "nan")
            & out["订单金额"].notna()
        ]

        if out.empty:
            continue

        out["月份"] = period
        out["来源文件"] = uploaded_file.name
        out["来源Sheet"] = sheet_name
        frames.append(out)

    if not frames:
        raise ValueError("未找到有效数据。请确认表格中存在“下单帐号”和“订单金额”列。")

    return pd.concat(frames, ignore_index=True), period


def build_comparison(raw: pd.DataFrame):
    """
    按 月份 + 下单帐号 汇总，并与上一自然月对比。
    """
    monthly = (
        raw.groupby(["月份", "下单帐号"], as_index=False)["订单金额"]
        .sum()
        .rename(columns={"订单金额": "本月下单金额"})
        .sort_values(["月份", "下单帐号"])
    )

    available_periods = set(monthly["月份"].unique())

    current = monthly.copy()
    current["上月月份"] = current["月份"] - 1

    prev = monthly[["下单帐号", "月份", "本月下单金额"]].rename(
        columns={
            "月份": "上月月份",
            "本月下单金额": "上月下单金额",
        }
    )

    comp = current.merge(prev, on=["下单帐号", "上月月份"], how="left")
    comp["有上月数据"] = comp["上月月份"].isin(available_periods)

    # 如果上一自然月有上传数据，但该账号没有记录，则上月金额视为 0；
    # 如果上一自然月根本没有上传数据，则显示“无上月数据”，不误标新增。
    comp["上月下单金额"] = comp.apply(
        lambda r: (
            (0 if pd.isna(r["上月下单金额"]) else r["上月下单金额"])
            if r["有上月数据"]
            else None
        ),
        axis=1,
    )

    comp["变化金额"] = comp.apply(
        lambda r: (
            r["本月下单金额"] - r["上月下单金额"]
            if r["有上月数据"]
            else None
        ),
        axis=1,
    )

    def get_status(r):
        if not r["有上月数据"]:
            return "无上月数据"

        prev_amt = r["上月下单金额"]
        if pd.isna(prev_amt):
            prev_amt = 0

        if prev_amt == 0:
            return "新增"

        diff = r["本月下单金额"] - prev_amt
        if diff > 0:
            return "增加"
        if diff < 0:
            return "减少"
        return "持平"

    comp["状态"] = comp.apply(get_status, axis=1)

    comp["月份"] = comp["月份"].astype(str)
    comp["上月月份"] = comp["上月月份"].astype(str)

    out_cols = [
        "月份",
        "下单帐号",
        "本月下单金额",
        "上月下单金额",
        "变化金额",
        "状态",
    ]
    comp = comp[out_cols]

    summary = (
        comp.groupby("月份", as_index=False)
        .agg(
            下单账号数=("下单帐号", "nunique"),
            总下单金额=("本月下单金额", "sum"),
            新增账号数=("状态", lambda s: (s == "新增").sum()),
        )
        .sort_values("月份")
    )

    monthly_for_pivot = monthly.copy()
    monthly_for_pivot["月份"] = monthly_for_pivot["月份"].astype(str)

    pivot = monthly_for_pivot.pivot_table(
        index="下单帐号",
        columns="月份",
        values="本月下单金额",
        aggfunc="sum",
        fill_value=0,
    ).reset_index()

    return comp, summary, pivot


def style_comparison_row(row):
    status = row["状态"]
    if status == "新增":
        color = "#FFF2CC"  # 黄色：新增账号
    elif status == "增加":
        color = "#E2EFDA"  # 浅绿
    elif status == "减少":
        color = "#FCE4D6"  # 浅红
    elif status == "无上月数据":
        color = "#F2F2F2"  # 浅灰
    else:
        color = ""

    return [f"background-color: {color}" if color else "" for _ in row]


def to_excel_bytes(comp: pd.DataFrame, summary: pd.DataFrame, pivot: pd.DataFrame):
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        comp.to_excel(writer, sheet_name="账号月度对比", index=False)
        summary.to_excel(writer, sheet_name="月度汇总", index=False)
        pivot.to_excel(writer, sheet_name="账号月份透视", index=False)

        wb = writer.book
        ws = wb["账号月度对比"]

        headers = {cell.value: cell.column for cell in ws[1]}
        status_col = headers["状态"]

        fills = {
            "新增": PatternFill(
                fill_type="solid", start_color="FFF2CC", end_color="FFF2CC"
            ),
            "增加": PatternFill(
                fill_type="solid", start_color="E2EFDA", end_color="E2EFDA"
            ),
            "减少": PatternFill(
                fill_type="solid", start_color="FCE4D6", end_color="FCE4D6"
            ),
            "无上月数据": PatternFill(
                fill_type="solid", start_color="F2F2F2", end_color="F2F2F2"
            ),
        }

        for row in range(2, ws.max_row + 1):
            status = ws.cell(row=row, column=status_col).value
            fill = fills.get(status)
            if fill:
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = fill

        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

        for col_name in ["本月下单金额", "上月下单金额", "变化金额"]:
            c = headers[col_name]
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=c).number_format = "#,##0.00"

        ws_summary = wb["月度汇总"]
        for col in range(1, ws_summary.max_column + 1):
            ws_summary.column_dimensions[get_column_letter(col)].width = 18

    output.seek(0)
    return output.getvalue()


st.title("📊 万商月度下单金额对比工具")
st.caption(
    "上传多个月份的万商表格，自动按“下单帐号 + 月份”汇总订单金额，"
    "并与上一自然月对比。新增账号使用黄色色块标记。"
)

with st.sidebar:
    st.header("参数设置")
    default_year = st.number_input(
        "文件名只有“6月/7月”时，使用的年份",
        min_value=2000,
        max_value=2100,
        value=datetime.date.today().year,
        step=1,
        help="例如文件名是“6月万商.xls”，则按此年份构造 6 月。"
    )

    st.markdown(
        """
        **文件名建议：**
        - `6月万商.xls`
        - `7月万商.xls`
        - `2025年6月万商.xls`
        - `2025-06万商.xlsx`
        """
    )

uploaded_files = st.file_uploader(
    "上传 Excel 文件，可一次选择多个月份",
    type=["xls", "xlsx"],
    accept_multiple_files=True,
)

if uploaded_files:
    all_rows = []
    detected = []

    for f in uploaded_files:
        try:
            df, period = read_one_file(f, int(default_year))
            all_rows.append(df)
            detected.append(f"{f.name} → {period}")
        except Exception as e:
            st.error(f"{f.name} 读取失败：{e}")

    if detected:
        st.success("已识别月份：\n\n" + "\n\n".join(detected))

    if all_rows:
        raw = pd.concat(all_rows, ignore_index=True)

        comp, summary, pivot = build_comparison(raw)

        st.subheader("1. 账号月度对比明细")
        st.caption(
            "本月下单金额 = 该账号当月“订单金额”合计；"
            "上月下单金额 = 该账号上一自然月“订单金额”合计。"
            "黄色 = 新增账号；浅绿 = 增加；浅红 = 减少；灰色 = 未上传上月数据。"
        )

        styled = (
            comp.style.apply(style_comparison_row, axis=1)
            .format(
                {
                    "本月下单金额": "{:,.2f}",
                    "上月下单金额": "{:,.2f}",
                    "变化金额": "{:,.2f}",
                },
                na_rep="",
            )
        )

        st.dataframe(styled, use_container_width=True, height=520)

        st.subheader("2. 月度汇总")
        st.dataframe(
            summary.style.format(
                {
                    "总下单金额": "{:,.2f}",
                }
            ),
            use_container_width=True,
        )

        st.subheader("3. 账号 × 月份 金额透视表")
        st.dataframe(pivot, use_container_width=True, height=420)

        excel_bytes = to_excel_bytes(comp, summary, pivot)

        st.download_button(
            label="⬇️ 下载带颜色标记的 Excel 结果",
            data=excel_bytes,
            file_name="万商月度下单金额对比.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        st.info(
            "说明：只有上传了上一自然月的数据，才会进行“新增/增加/减少”判断。"
            "例如只上传 7 月，不上传 6 月，则 7 月所有账号都会显示“无上月数据”，"
            "不会误标为新增。上传 6 月和 7 月后，7 月才能正确对比 6 月。"
        )
else:
    st.info("请先上传 Excel 文件。")